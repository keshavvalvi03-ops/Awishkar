"""
ClauseGuard AI: Web Application Backend API
-------------------------------------------
FastAPI backend connecting the 6-layer Explainable AI Contract Risk Analyser
to the interactive web application.

Features:
- Full 6-layer contract risk analysis pipeline execution
- Live contract type suggestion with evidence and uncertainty handling
- Genuine file extraction for TXT, PDF, and DOCX documents
- Preloaded sample contracts across all 10 supported contract types
- Offline static file serving for modern single-page dashboard
"""

import os
import io
from typing import Dict, List, Optional, Any
from fastapi import FastAPI, UploadFile, File, HTTPException, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse, JSONResponse
from pydantic import BaseModel, Field

from pipeline import full_analysis, load_cuad_questions
from contract_types import (
    CONTRACT_TYPES,
    list_supported_contract_types,
    suggest_contract_type,
    get_contract_type,
    STATUS_DETECTED,
    STATUS_NOT_DETECTED,
    STATUS_UNCERTAIN,
    RELEVANCE_EXPECTED,
    RELEVANCE_CONDITIONAL,
    RELEVANCE_OPTIONAL,
)
from custom_detector import scan_all_clauses, detect_clause_in_text
from demo_multi_contract import SAMPLE_CONTRACTS
from demo import sample_extractions as baseline_extractions, sample_raw_contract_text as baseline_text

# Initialize FastAPI app
app = FastAPI(
    title="ClauseGuard AI API",
    description="Explainable AI-Based Clause Risk Analyser in Contract",
    version="2.0.0",
)

# Enable CORS for local development and integration
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

STATIC_DIR = os.path.join(os.path.dirname(__file__), "static")


# Pydantic Request Models
class AnalyzeRequest(BaseModel):
    text: Optional[str] = Field(default="", description="Raw contract text to analyse")
    contract_type: Optional[str] = Field(default=None, description="Selected contract type name or alias")
    extractions: Optional[List[Dict[str, Any]]] = Field(default=None, description="Optional pre-extracted clause spans")
    is_sample: Optional[bool] = Field(default=False, description="Flag indicating sample demonstration data")


class SuggestTypeRequest(BaseModel):
    text: str = Field(..., description="Contract text for heuristic type suggestion")


# API Endpoints
@app.get("/api/health")
def health_check():
    """Health check endpoint confirming API status and architecture layers."""
    return {
        "status": "healthy",
        "app": "ClauseGuard AI",
        "project": "Explainable AI-Based Clause Risk Analyser in Contract",
        "version": "2.0.0",
        "layers_available": 6,
        "supported_contract_types": len(CONTRACT_TYPES),
    }


@app.get("/api/contract-types")
def get_contract_types_list():
    """Returns metadata and checklists for all 10 supported contract types."""
    types_data = []
    for ct_id, ct in CONTRACT_TYPES.items():
        types_data.append({
            "type_id": ct.type_id,
            "name": ct.name,
            "description": ct.description,
            "aliases": ct.aliases,
            "checklist_count": len(ct.checklist),
            "expected_count": sum(1 for item in ct.checklist if item.relevance == RELEVANCE_EXPECTED),
            "conditional_count": sum(1 for item in ct.checklist if item.relevance == RELEVANCE_CONDITIONAL),
            "checklist": [item.to_dict() for item in ct.checklist],
        })
    return {
        "contract_types": types_data,
        "count": len(types_data),
    }


@app.post("/api/suggest-type")
def suggest_type_endpoint(req: SuggestTypeRequest):
    """Suggests the most likely contract type with evidence snippets and confidence."""
    if not req.text or not req.text.strip():
        raise HTTPException(status_code=400, detail="Contract text is required for suggestion.")
    suggestion = suggest_contract_type(req.text)
    return suggestion


@app.get("/api/samples")
def get_sample_contracts():
    """
    Returns pre-configured sample contracts for interactive demonstration:
    - The baseline demonstration from demo.py (score 8.16/10 High)
    - 10 contract types from demo_multi_contract.py
    """
    samples = {
        "Baseline Demo (Original PoC)": {
            "contract_type": "General Business Agreement",
            "description": "Baseline demonstration with Uncapped Liability, Cap On Liability, and convenience termination.",
            "text": baseline_text.strip(),
            "extractions": baseline_extractions,
            "expected_score": 8.16,
            "expected_level": "High",
        }
    }

    for ct_name, data in SAMPLE_CONTRACTS.items():
        samples[ct_name] = {
            "contract_type": ct_name,
            "description": f"Standard commercial sample representing {ct_name}.",
            "text": data["text"].strip(),
            "extractions": data["extractions"],
        }

    return {"samples": samples}


@app.post("/api/upload")
async def upload_document(file: UploadFile = File(...)):
    """
    Extracts text from uploaded .txt, .pdf, or .docx files.
    Rejects unsupported files with transparent error reporting.
    """
    filename = file.filename or "unknown"
    ext = os.path.splitext(filename)[1].lower()

    content_bytes = await file.read()
    if not content_bytes:
        raise HTTPException(status_code=400, detail="The uploaded file is empty.")

    extracted_text = ""

    # 1. Plain text (.txt)
    if ext == ".txt":
        try:
            extracted_text = content_bytes.decode("utf-8")
        except UnicodeDecodeError:
            extracted_text = content_bytes.decode("latin-1", errors="replace")

    # 2. Microsoft Word (.docx)
    elif ext == ".docx":
        try:
            import docx
            doc = docx.Document(io.BytesIO(content_bytes))
            full_paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        full_paragraphs.append(row_text)
            extracted_text = "\n\n".join(full_paragraphs)
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Failed to extract text from DOCX: {str(e)}")

    # 3. PDF (.pdf)
    elif ext == ".pdf":
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(content_bytes))
            pages_text = []
            for i, page in enumerate(reader.pages):
                page_content = page.extract_text()
                if page_content:
                    pages_text.append(page_content.strip())
            extracted_text = "\n\n--- Page Break ---\n\n".join(pages_text)
            if not extracted_text.strip():
                raise HTTPException(
                    status_code=400,
                    detail="No extractable text was found in this PDF (it may contain scanned image pages without text OCR)."
                )
        except Exception as e:
            if isinstance(e, HTTPException):
                raise e
            raise HTTPException(status_code=400, detail=f"Failed to extract text from PDF: {str(e)}")

    else:
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported file format '{ext}'. Only .txt, .pdf, and .docx documents are supported."
        )

    # Suggest contract type based on extracted text
    suggestion = suggest_contract_type(extracted_text)

    return {
        "filename": filename,
        "extension": ext,
        "char_count": len(extracted_text),
        "word_count": len(extracted_text.split()),
        "text": extracted_text,
        "suggestion": suggestion,
    }


@app.post("/api/analyze")
def analyze_contract_endpoint(req: AnalyzeRequest):
    """
    Core contract risk analysis endpoint executing the complete 6-layer pipeline.
    """
    raw_text = req.text.strip() if req.text else ""
    extractions = req.extractions or []

    # Validate input: either raw text or extractions must be provided
    if not raw_text and not extractions:
        raise HTTPException(
            status_code=400,
            detail="Contract text or structured extractions must be provided for analysis."
        )

    # If raw text is provided without pre-extracted spans, run the rule-based custom detector
    if raw_text and not extractions:
        detected_spans = scan_all_clauses(raw_text)
        extractions = [
            {
                "category": d["category"],
                "text": d["text"],
                "confidence": d["confidence"],
                "detection_method": d["detection_method"],
            }
            for d in detected_spans
        ]

    # Determine contract type
    selected_type = req.contract_type
    type_suggestion = None
    if raw_text:
        type_suggestion = suggest_contract_type(raw_text, extractions)
        if not selected_type:
            selected_type = type_suggestion.get("primary_suggestion")

    # Run complete 6-layer analysis
    try:
        report = full_analysis(
            extractions=extractions,
            raw_contract_text=raw_text if raw_text else None,
            contract_type=selected_type,
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Analysis pipeline error: {str(e)}")

    # Calculate summary metrics
    clauses = report.get("clauses", [])
    high_risk_count = sum(1 for c in clauses if c.get("risk_level") == "High")
    med_risk_count = sum(1 for c in clauses if c.get("risk_level") == "Medium")
    low_risk_count = sum(1 for c in clauses if c.get("risk_level") == "Low")
    na_count = sum(1 for c in clauses if c.get("risk_level") == "Not Applicable")

    checklist = report.get("contract_checklist", [])
    detected_checklist = sum(1 for item in checklist if item.get("status") == STATUS_DETECTED)
    missing_expected = sum(
        1 for item in checklist
        if item.get("status") == STATUS_NOT_DETECTED and item.get("relevance") == RELEVANCE_EXPECTED
    )
    missing_conditional = sum(
        1 for item in checklist
        if item.get("status") == STATUS_NOT_DETECTED and item.get("relevance") == RELEVANCE_CONDITIONAL
    )

    summary_metrics = {
        "total_extracted_clauses": len(clauses),
        "risk_bearing_clauses": high_risk_count + med_risk_count + low_risk_count,
        "high_risk_clauses": high_risk_count,
        "medium_risk_clauses": med_risk_count,
        "low_risk_clauses": low_risk_count,
        "not_applicable_clauses": na_count,
        "checklist_total": len(checklist),
        "checklist_detected": detected_checklist,
        "missing_expected_clauses": missing_expected,
        "missing_conditional_clauses": missing_conditional,
        "is_sample_data": bool(req.is_sample),
    }

    # Build response payload
    return {
        "contract_type": report.get("contract_type"),
        "suggestion": type_suggestion,
        "contract_risk": report.get("contract_risk"),
        "summary": summary_metrics,
        "clauses": clauses,
        "missing_clauses": report.get("missing_clauses", []),
        "contract_checklist": checklist,
    }


# Mount Static Files and Serve Frontend SPA
if os.path.exists(STATIC_DIR):
    app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")

    @app.get("/")
    def serve_frontend():
        index_path = os.path.join(STATIC_DIR, "index.html")
        if os.path.exists(index_path):
            return FileResponse(index_path)
        return {"message": "ClauseGuard AI backend running. Place index.html in static/"}

    @app.get("/favicon.ico", include_in_schema=False)
    def favicon():
        fav_path = os.path.join(STATIC_DIR, "favicon.svg")
        if os.path.exists(fav_path):
            return FileResponse(fav_path, media_type="image/svg+xml")
        return JSONResponse(status_code=404, content={"message": "Favicon not found"})


if __name__ == "__main__":
    import uvicorn
    print("\n" + "=" * 70)
    print("Starting ClauseGuard AI Web Server on http://127.0.0.1:8000")
    print("=" * 70 + "\n")
    uvicorn.run("app:app", host="127.0.0.1", port=8000, reload=True)
